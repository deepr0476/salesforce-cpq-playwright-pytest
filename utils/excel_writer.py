from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from utils.system_logger import get_logger

logger = get_logger(__name__)

RESULT_HEADERS = [
    "test_case",
    "status",
    "account_name",
    "account_id",
    "contact_name",
    "contact_id",
    "opportunity_id",
    "quote_id",
    "order_id",
    "contract_id",
    "error_message",
    "execution_timestamp",
    "duration_seconds",
]


class ExcelWriter:
    def __init__(self, path: Path):
        self.path = path
        self.path.parent.mkdir(parents=True, exist_ok=True)

    def reset_execution_results(self, sheet: str = "execution_results") -> None:
        """Clears old historical run records and rewrites fresh headers."""
        if not self.path.exists():
            return
        workbook = load_workbook(self.path)
        if sheet in workbook.sheetnames:
            del workbook[sheet]
        ws = workbook.create_sheet(sheet)
        ws.append(RESULT_HEADERS)
        workbook.save(self.path)
        logger.info("STATUS=PASS ACTION=RESET_RESULTS_SHEET SHEET=%s", sheet)

    def write_result(self, sheet: str, data: dict[str, Any]) -> None:
        if self.path.exists():
            workbook = load_workbook(self.path)
        else:
            workbook = Workbook()
            default = workbook.active
            workbook.remove(default)

        if sheet not in workbook.sheetnames:
            ws = workbook.create_sheet(sheet)
            ws.append(RESULT_HEADERS)
        else:
            ws = workbook[sheet]
            if ws.max_row == 0 or not ws.cell(1, 1).value:
                ws.append(RESULT_HEADERS)

        headers = [cell.value for cell in ws[1]]
        row = [data.get(h, "") for h in headers]
        ws.append(row)
        workbook.save(self.path)
        logger.info("STATUS=PASS ACTION=WRITE_RESULT SHEET=%s TEST=%s", sheet, data.get("test_case"))


def now_ts() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")
