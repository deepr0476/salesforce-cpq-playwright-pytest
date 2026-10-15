from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from utils.system_logger import get_logger

logger = get_logger(__name__)


def _norm(value: Any) -> str:
    return str(value).strip() if value is not None else ""


class ExcelReader:
    def __init__(self, path: Path):
        self.path = path
        if not path.exists():
            raise FileNotFoundError(f"Excel workbook not found: {path}")
        self._workbook = load_workbook(path, data_only=True)

    def _sheet(self, name: str):
        if name not in self._workbook.sheetnames:
            raise ValueError(f"Sheet '{name}' not found. Available: {self._workbook.sheetnames}")
        return self._workbook[name]

    def get_records(self, sheet: str) -> list[dict[str, Any]]:
        ws = self._sheet(sheet)
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            raise ValueError(f"Sheet '{sheet}' is empty")
        headers = [_norm(h) for h in rows[0]]
        if not any(headers):
            raise ValueError(f"Sheet '{sheet}' has no headers")
        records: list[dict[str, Any]] = []
        for row in rows[1:]:
            if all(cell is None or str(cell).strip() == "" for cell in row):
                continue
            record = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers)) if headers[i]}
            records.append(record)
        logger.info("STATUS=PASS ACTION=READ_SHEET SHEET=%s ROWS=%s", sheet, len(records))
        return records

    def get_first_row(self, sheet: str) -> dict[str, Any]:
        records = self.get_records(sheet)
        if not records:
            raise ValueError(f"Sheet '{sheet}' has no data rows")
        return records[0]

    def get_row(self, sheet: str, key: str, value: Any) -> dict[str, Any]:
        records = self.get_records(sheet)
        for record in records:
            if str(record.get(key, "")).strip() == str(value).strip():
                return record
        raise ValueError(f"No row in '{sheet}' where {key}={value}")
